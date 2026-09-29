from datetime import datetime, timezone

import openpyxl

from dora_metrics.excel_report_builder import (
    SHEET_JOBS,
    SHEET_LEAD_TIME,
    SHEET_PARAMETERS,
    SHEET_REPOS,
    SHEET_RUNS,
    SHEET_SUMMARY,
    ExcelReportBuilder,
)
from dora_metrics.models import (
    JobInfo,
    LeadTimeRecord,
    ReportData,
    ReportParameters,
    RepoSummary,
    RunnerUsageStat,
    WeeklyMetric,
    WorkflowRun,
)


def _dt(iso: str) -> datetime:
    return datetime.fromisoformat(iso).replace(tzinfo=timezone.utc)


def _build_report_data() -> ReportData:
    weekly_metrics = [
        WeeklyMetric(
            iso_week="2024-W12",
            repo="lib-ios-core",
            repo_type="ios",
            deployment_frequency=3,
            lead_time_median_hours=12.5,
            lead_time_p90_hours=20.0,
            change_failure_rate_pct=10.0,
            mttr_hours=4.0,
        )
    ]
    repo_summaries = [
        RepoSummary(
            repo="lib-ios-core",
            repo_type="ios",
            visibility="private",
            default_branch="main",
            workflows_detected=2,
            last_run_at=_dt("2024-03-18T10:00:00"),
        )
    ]
    workflow_runs = [
        WorkflowRun(
            repo="lib-ios-core",
            workflow_id=1,
            workflow_name="Release iOS",
            workflow_path="release-ios.yml",
            run_id=1,
            name="Release iOS",
            event="push",
            head_branch="main",
            head_sha="abc123",
            status="completed",
            conclusion="success",
            run_attempt=1,
            created_at=_dt("2024-03-18T10:00:00"),
            run_started_at=_dt("2024-03-18T10:01:00"),
            updated_at=_dt("2024-03-18T10:11:00"),
            actor="jdoe",
            is_deploy=True,
        )
    ]
    jobs = [
        JobInfo(
            repo="lib-ios-core",
            run_id=1,
            job_id=1,
            name="build",
            runner_name="GitHub Actions 1",
            runner_group_name="Default",
            labels=("ubuntu-latest",),
            created_at=_dt("2024-03-18T10:00:00"),
            started_at=_dt("2024-03-18T10:01:00"),
            completed_at=_dt("2024-03-18T10:05:00"),
        )
    ]
    lead_time_records = [
        LeadTimeRecord(
            repo="lib-ios-core",
            pr_or_commit="#42",
            merged_at=_dt("2024-03-17T08:00:00"),
            deploy_run_id=1,
            deploy_completed_at=_dt("2024-03-18T10:11:00"),
            lead_time_hours=26.18,
        )
    ]
    runner_usage_stats = [
        RunnerUsageStat(
            repo="lib-ios-core",
            runner_label="ubuntu-latest",
            hosted_type="github-hosted",
            sample_size=1,
            queue_median_seconds=60.0,
            queue_p90_seconds=60.0,
            execution_median_seconds=240.0,
            execution_p90_seconds=240.0,
        )
    ]
    parameters = ReportParameters(
        org="my-org",
        repo_filter_type="name_pattern",
        repo_filter_value="^lib-",
        range_start="2024-03-18",
        range_end="2024-03-24",
        generated_at=_dt("2024-03-25T09:00:00"),
        triggered_by="jdoe",
    )

    return ReportData(
        weekly_metrics=weekly_metrics,
        repo_summaries=repo_summaries,
        workflow_runs=workflow_runs,
        jobs=jobs,
        lead_time_records=lead_time_records,
        runner_usage_stats=runner_usage_stats,
        parameters=parameters,
    )


def test_build_creates_six_sheets_with_expected_names(tmp_path):
    output_path = tmp_path / "dora_metrics_2024-03-18_2024-03-24.xlsx"
    builder = ExcelReportBuilder()

    result_path = builder.build(_build_report_data(), output_path)

    assert result_path.exists()
    workbook = openpyxl.load_workbook(result_path)
    assert workbook.sheetnames == [
        SHEET_SUMMARY,
        SHEET_REPOS,
        SHEET_RUNS,
        SHEET_JOBS,
        SHEET_LEAD_TIME,
        SHEET_PARAMETERS,
    ]


def test_summary_sheet_has_expected_columns_and_values(tmp_path):
    output_path = tmp_path / "report.xlsx"
    builder = ExcelReportBuilder()
    result_path = builder.build(_build_report_data(), output_path)

    workbook = openpyxl.load_workbook(result_path)
    sheet = workbook[SHEET_SUMMARY]
    header = [cell.value for cell in sheet[1]]
    assert header == [
        "Semana",
        "Repo",
        "Tipo",
        "Deployment Frequency",
        "Lead Time mediana (h)",
        "Lead Time p90 (h)",
        "Change Failure Rate (%)",
        "MTTR (h)",
    ]
    data_row = [cell.value for cell in sheet[2]]
    assert data_row[0] == "2024-W12"
    assert data_row[1] == "lib-ios-core"
    assert data_row[3] == 3


def test_parameters_sheet_contains_org_and_filter(tmp_path):
    output_path = tmp_path / "report.xlsx"
    builder = ExcelReportBuilder()
    result_path = builder.build(_build_report_data(), output_path)

    workbook = openpyxl.load_workbook(result_path)
    sheet = workbook[SHEET_PARAMETERS]
    values = {row[0].value: row[1].value for row in sheet.iter_rows(min_row=2)}
    assert values["Organizacion"] == "my-org"
    assert values["Valor de filtro"] == "^lib-"


def test_jobs_sheet_includes_raw_jobs_and_runner_usage_summary(tmp_path):
    output_path = tmp_path / "report.xlsx"
    builder = ExcelReportBuilder()
    result_path = builder.build(_build_report_data(), output_path)

    workbook = openpyxl.load_workbook(result_path)
    sheet = workbook[SHEET_JOBS]
    all_values = [cell.value for row in sheet.iter_rows() for cell in row]
    assert "build" in all_values
    assert "Resumen de uso de runners" in all_values
